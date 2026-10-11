# Builds manual-qa/Signal-School-Test-Cases.xlsx (needs: pip install openpyxl).
#   python3 scripts/manual-qa/test_cases.py
# The children and logins named here come from the API's QA seed (`npm run db:seed:qa`, scripts/demo/qa.js there):
# change both together.
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[2] / 'manual-qa' / 'Signal-School-Test-Cases.xlsx'

RED = 'C8102E'
STATUSES = ['Not run', 'Pass', 'Fail', 'Blocked', 'Skipped']

# Logins (all from the QA data).
LOGINS = [
    ('Teacher', 'Sunita Patil', 'sunita@demo.test', '9820000004', 'password123', 'Class teacher of Std 1 A. Std 1 A attendance is not taken today.'),
    ('Teacher', 'Rahul Verma', 'rahul@demo.test', '9820000005', 'password123', 'Class teacher of Std 5 A; teaches Maths and English in Std 5, 6, 7.'),
    ('Teacher', 'Kavita Jadhav', 'kavita@demo.test', '9820000006', 'password123', 'Class teacher of Std 2 A and Std 3 A. Std 3 A attendance is not taken today.'),
    ('New teacher', 'Neha Joshi', 'newteacher@demo.test', '9820000011', 'Welcome@123 (temporary)', 'Logs in for the first time: must choose a new password. No class yet.'),
    ('Inactive teacher', 'Deepak Rane', 'inactive@demo.test', '9820000012', 'password123', 'Made inactive: must NOT be able to log in.'),
    ('Office clerk', 'Ganesh Pawar', 'clerk@demo.test', '9820000003', 'password123', 'Admits children, imports, health, leaving certificates. Cannot take attendance.'),
    ('Principal', 'Anita Deshmukh', 'principal@demo.test', '9820000002', 'password123', 'Runs Signal Shala, Thane: everything in that school.'),
    ('Trust owner', 'Meera Kulkarni', 'owner@demo.test', '9820000001', 'password123', 'Sees both schools (Thane and Nerul). Can add schools.'),
    ('Teacher (2nd school)', 'Imran Shaikh', 'imran@demo.test', '9820000010', 'password123', 'Teacher at Signal Shala, Nerul (that school uses the "Marigold" look).'),
]

# Fixed children for testing (Signal Shala, Thane, current year).
CHILDREN = [
    ('Riya Kale', 'Std 1 A', 'Girl, has a photo. Parent Sunil Kale, 9123400001. Was in Balwadi A last year.', 'Search, profile, attendance, diary, health, past years'),
    ('Aryan Mhatre', 'Std 1 A', 'Boy, NO photo. Parent Kavita Mhatre, 9123400002.', 'Adding a photo'),
    ('Ishaan Jagtap', 'Std 2 A', 'Boy. His birthday is TODAY.', 'Birthday line on the Dashboard'),
    ('Swati Gholap', 'Std 2 A', 'Girl. Last health check-up: "needs a doctor again" (eyes).', 'Health follow-up'),
    ('Nikhil Sonawane', 'Std 3 A', 'Boy, studying.', 'Mark as left school + leaving certificate'),
    ('Kiran Thakur', 'Std 4 A', 'Girl. Absent the last 5 school days, low attendance.', 'Dashboard: "Absent several days in a row" and "Children who often miss school"'),
    ('Lata Bhalerao', 'Std 2 A (last year)', 'Girl. Left school last year (family moved away).', '"Search all years" and Re-admit'),
]

FILES = [
    ('child-photo-girl.jpg', 'A drawn photo of a girl', 'Add / change a child\'s photo'),
    ('child-photo-boy.jpg', 'A drawn photo of a boy', 'Add a photo for Aryan Mhatre'),
    ('teacher-photo.jpg', 'A drawn photo of a woman', 'Your own photo on the "Me" page'),
    ('school-logo.png', 'A sample school logo', 'School settings: Upload school logo'),
    ('photo-too-big-9MB.png', 'A picture larger than 8 MB', 'Must be refused: "The file is too large"'),
    ('not-a-photo.pdf', 'A PDF, not a picture or sheet', 'Must be refused: "This file type is not allowed"'),
    ('students-import-5-good.xlsx', '5 correct children', 'Import: all 5 rows ready'),
    ('students-import-with-errors.xlsx', '5 rows: 1 correct, 4 with mistakes', 'Import: 1 ready, 4 show problems'),
    ('students-import-3-rows.csv', '3 correct children (CSV file)', 'Import from a CSV file'),
]

# (area, login, device, where, steps, test data, expected, priority)
P = 'Phone'
C = 'Computer'
B = 'Phone + Computer'
CASES = [
    # ---- Login
    ('Login', '—', P, 'Open the app link', 'Open the link for the first time. Tap "मराठी", then "हिंदी", then "ગુજરાતી", then "English".', '', 'The login screen changes language each time you tap. Logo, "Signal School" and the Log in box are visible without zooming.', 'High'),
    ('Login', 'Teacher (email)', P, 'Login screen', 'Type the email and password, tap "Log in".', 'sunita@demo.test / password123', 'Opens the "Today" page with "Namaste, Sunita".', 'High'),
    ('Login', 'Teacher (mobile)', P, 'Login screen', 'Log out. Log in with the mobile number instead of email.', '9820000004 / password123', 'Opens the "Today" page with "Namaste, Sunita".', 'High'),
    ('Login', 'Teacher', P, 'Login screen', 'Type a wrong password and tap "Log in".', 'sunita@demo.test / wrong123', 'Red message "Wrong mobile number/email or password". You stay on the login screen.', 'High'),
    ('Login', '—', P, 'Login screen', 'Leave both boxes empty and tap "Log in".', '', 'You are not logged in. The empty boxes are marked as needed.', 'Medium'),
    ('Login', '—', P, 'Login screen', 'Type a password. Tap the eye icon, then tap it again.', 'any', 'The password shows as letters, then hides again as dots.', 'Low'),
    ('Login', 'Inactive teacher', P, 'Login screen', 'Log in with the inactive account.', 'inactive@demo.test / password123', 'Message "This account is inactive. Contact the school office." Not logged in.', 'High'),
    ('Login', 'New teacher', P, 'Login → Welcome page', 'Log in with the temporary password. On "Welcome, Neha Joshi" type current password, then a new password and a DIFFERENT one in "Type the new password again". Tap "Change password".', 'newteacher@demo.test / Welcome@123; new: NehaQa@2026; again: NehaQa@2027', 'Message "Passwords do not match". Password is not changed.', 'High'),
    ('Login', 'New teacher', P, 'Welcome page', 'Now type a short new password (both boxes the same) and tap "Change password".', 'new: abc12', 'Message "At least 8 characters". Password is not changed.', 'Medium'),
    ('Login', 'New teacher', P, 'Welcome page', 'Type the right values and tap "Change password".', 'current: Welcome@123; new and again: NehaQa@2026', '"Password changed". The Today page opens and says "No class is assigned to you yet". Log out and log in with NehaQa@2026: works.', 'High'),
    ('Login', 'Teacher', P, 'Me → Log out', 'Tap "More" → "Me" → "Log out" → confirm "Log out". Then press the phone\'s Back button.', '', 'The login screen opens. Back does not show the app pages again.', 'High'),
    ('Login', 'Teacher', P, 'Any page', 'Log in, close the browser tab completely, open the app link again.', '', 'You are still logged in (no need to type the password again).', 'Medium'),
    ('Login', 'Teacher', C, 'Login screen', 'Type a wrong password 10 times for Pooja (pooja@demo.test). Then type the right password.', 'pooja@demo.test / wrong… then password123', 'After 10 wrong tries: "Too many wrong attempts. Try again after 15 minutes." The right password also waits 15 minutes.', 'Low'),
    # ---- Teacher: today and attendance
    ('Teacher – Today', 'Teacher', P, 'Today', 'Log in as Sunita. Look at the Today page.', 'sunita@demo.test', 'Shows "Namaste, Sunita", today\'s date, "Attendance is waiting in 1 class", class card "Std 1 A – Attendance not taken yet" with a big "Take attendance" button, and "My subjects" with progress bars.', 'High'),
    ('Teacher – Today', 'Teacher', P, 'Bottom bar', 'Tap each item of the bottom bar: Today, Attendance, Students, Syllabus, More.', '', 'Each opens its page. "More" opens a small sheet with "Marks" and "Me". The bar stays at the bottom and never hides buttons.', 'High'),
    ('Teacher – Attendance', 'Teacher', P, 'Today → Take attendance', 'Tap "Take attendance" on Std 1 A.', '', 'List of all Std 1 A children with roll numbers. Everyone starts on P (green). Counts show "Present: (all) · Absent: 0 · Leave: 0". A short help note is shown.', 'High'),
    ('Teacher – Attendance', 'Teacher', P, 'Take attendance', 'Tap "A" for Aryan Mhatre and "L" for Riya Kale.', 'Aryan Mhatre, Riya Kale', 'Aryan\'s row turns red (A), Riya\'s turns yellow (L). Counts change at once: Absent 1, Leave 1.', 'High'),
    ('Teacher – Attendance', 'Teacher', P, 'Take attendance', 'Type a name in the "Search" box above the list. Then clear it.', 'Riya', 'Only Riya Kale is listed. After clearing, all children are back with your marks kept.', 'Medium'),
    ('Teacher – Attendance', 'Teacher', P, 'Take attendance', 'Tap "Save attendance".', '', 'Message "Attendance saved: … present, 1 absent". Today page now shows Std 1 A in green "Attendance done: … present (time)" and the button says "Change attendance".', 'High'),
    ('Teacher – Attendance', 'Teacher', P, 'Today → Change attendance', 'Open it again. Change Aryan to P. Do not save yet.', '', 'A note says "You changed some marks. Press Save attendance to keep them." After "Save attendance" the new counts are saved.', 'High'),
    ('Teacher – Attendance', 'Teacher', P, 'Take attendance', 'Mark 2 children A, then tap "Mark everyone present".', '', 'Everyone is back on P. Counts show Absent 0.', 'Medium'),
    ('Teacher – Attendance', 'Teacher', P, 'Take attendance → Date', 'Change the date to yesterday.', 'yesterday', 'Yesterday\'s saved attendance of Std 1 A is shown and can be changed and saved.', 'Medium'),
    ('Teacher – Attendance', 'Teacher', P, 'Take attendance → Date', 'Change the date to a day 10 days ago.', '10 days ago', 'Warning "This date is too old to change. Ask the principal." The buttons cannot be changed and there is no Save button.', 'High'),
    ('Teacher – Attendance', 'Teacher', P, 'Take attendance → Date', 'Try to pick tomorrow\'s date.', 'tomorrow', 'Not allowed: warning "You cannot choose a future date" (or the date cannot be picked).', 'Medium'),
    ('Teacher – Attendance', 'Teacher', P, 'Take attendance (no internet)', 'Turn ON Airplane mode. Change one mark and tap "Save attendance". Then turn Airplane mode OFF and wait 1 minute.', '', 'With no internet: "No internet. Attendance is saved on this phone and will be sent automatically…". After internet returns, the waiting message goes away and the change is saved (check it on a computer).', 'High'),
    ('Teacher – Diary', 'Teacher', P, 'Today → Class diary', 'Tap "Class diary" on Std 1 A. Write a note and save.', 'Today we learned numbers 1 to 20.', '"Diary saved". The note shows with today\'s date and "by Sunita Patil".', 'Medium'),
    ('Teacher – Students', 'Teacher', P, 'Today → Students', 'Tap "Students" on the Std 1 A card.', '', 'The Std 1 A class list opens with photos (or initials), roll numbers and today\'s mark.', 'Medium'),
    ('Teacher – Students', 'Teacher', P, 'Students → search', 'Type in the search box.', 'Riya', 'Riya Kale appears (Std 1 A).', 'High'),
    ('Teacher – Students', 'Teacher', P, 'Students → search', 'Search by the parent\'s mobile number.', '9123400004', 'Kiran Thakur appears.', 'Medium'),
    ('Teacher – Student profile', 'Teacher', P, 'Students → Riya Kale', 'Open Riya Kale.', '', 'Photo, "Std 1 A", Roll no., GR number, age, "Sunil Kale · Father" with call and WhatsApp icons, and tabs: Details, Daily diary, Past years, Report card, Health. "Attendance this month" shows a small calendar with totals.', 'High'),
    ('Teacher – Student profile', 'Teacher', P, 'Riya Kale → Past years', 'Open the "Past years" tab.', '', 'Shows last year: Balwadi A, "Moved to next class", with that year\'s attendance.', 'Medium'),
    ('Teacher – Student profile', 'Teacher', P, 'Riya Kale → Daily diary', 'Add a note about the child and save.', 'Riya read a full story today.', 'Note appears at the top with today\'s date and "by Sunita Patil".', 'Medium'),
    ('Teacher – Student profile', 'Teacher', P, 'Riya Kale → call / WhatsApp', 'Tap the phone icon. Go back. Tap the WhatsApp icon.', '', 'Phone dialer opens with 9123400001. WhatsApp opens a chat to that number (if WhatsApp is installed).', 'Medium'),
    ('Teacher – Student profile', 'Teacher', P, 'Aryan Mhatre → camera icon', 'Open Aryan Mhatre (no photo). Tap the camera icon on the picture, choose the photo file.', 'test-files/child-photo-boy.jpg', 'The photo shows on the profile and in the class list.', 'High'),
    ('Teacher – Student profile', 'Teacher', P, 'Aryan Mhatre → camera icon', 'Try to upload the wrong files as a photo.', 'not-a-photo.pdf, then photo-too-big-9MB.png', 'PDF: "This file type is not allowed" (or it cannot be picked). Big file: "The file is too large". The old photo stays.', 'Medium'),
    ('Teacher – Student profile', 'Teacher', P, 'Riya Kale → Edit', 'Tap "Edit", change the address, tap "Save student".', 'New address: Kopri, Thane', '"Saved". Details tab shows the new address.', 'Medium'),
    ('Teacher – Student profile', 'Teacher', P, 'Riya Kale → Health', 'Add a health check-up: today\'s date, height and weight. Save.', 'Height 112, Weight 19', 'The check-up shows in the list with "Height: 112 cm", "Weight: 19 kg".', 'Medium'),
    ('Teacher – Students', 'Teacher', P, 'Students → Add student', 'Tap "Add student". Fill name, class, gender. Tap "Save student".', 'QA Child One, Std 1 A, Girl', 'Message "QA Child One added with register number SS-…". The child is in the Std 1 A list.', 'High'),
    ('Teacher – Students', 'Teacher', P, 'Students → Add student', 'Add a child with the same name and parent mobile as an existing child.', 'Riya Kale, Std 1 A, mobile 9123400001', 'Question "Is this child already in the school records?" shows Riya Kale. Tap Cancel: nothing is added.', 'High'),
    ('Teacher – Syllabus', 'Teacher', P, 'Syllabus → Maths · Std 1 A', 'Open a subject. Tap "Mark as taught" on a topic that is not taught yet. Then undo with "Mark as not taught".', 'Maths · Std 1 A', 'The topic shows "Taught on (today) by Me" and the % goes up; after undo it goes back.', 'High'),
    ('Teacher – Marks', 'Teacher', P, 'More → Marks', 'Choose Class Std 1 A, Subject Maths, Term Semester 1. In "Maximum marks for everyone" type 50 and tap "Fill in". Enter marks for 3 children. Save.', 'Marks: 45, 30, 12', '"Saved". Reopen Marks: the numbers are still there.', 'High'),
    ('Teacher – Marks', 'Teacher', P, 'More → Marks', 'Enter marks higher than "Out of".', '60 out of 50', 'Message "More than maximum marks". Not saved.', 'Medium'),
    ('Teacher – Marks', 'Teacher', P, 'Riya Kale → Report card', 'Open the "Report card" tab, choose Semester 1, tap "Open report card", then "Print / Save PDF".', '', 'Shows school name, child details, marks, attendance line and signature lines (Class teacher, Principal, Parent). The print / save as PDF screen opens.', 'Medium'),
    ('Teacher – Old year', 'Teacher', P, 'Year box at the top', 'Change the year at the top from 2026-27 to the year before. Open Attendance and a child. Then choose the current year again.', '', 'Banner "You are viewing … Old years can be seen but not changed." Nothing can be saved in the old year. "Back to …" returns to the current year.', 'Medium'),
    ('Teacher – Me', 'Teacher', P, 'More → Me → Language', 'Choose हिंदी. Open Today, Attendance, Students. Then choose English again.', '', 'The whole app (menus, buttons, messages) is in Hindi, then back to English. Note any word still in English.', 'High'),
    ('Teacher – Me', 'Teacher', P, 'More → Me → Text size', 'Choose the biggest "Aa". Open Today and Take attendance. Then choose Normal again.', '', 'Text gets bigger everywhere. Nothing overlaps or is cut, no sideways scrolling. Back to normal after.', 'Medium'),
    ('Teacher – Me', 'Teacher', P, 'More → Me → photo', 'Tap the camera on your own picture and choose a photo.', 'test-files/teacher-photo.jpg', 'Your photo shows on the Me page.', 'Low'),
    ('Teacher – Access', 'Teacher', P, 'Menus', 'Look at the menus. Then type /staff at the end of the app address and open it.', '', 'No "Teachers & staff", "School settings", "Academic years" or "Activity log" for a teacher. Opening /staff goes back to the Today page.', 'High'),
    ('Teacher – Subject teacher', 'Teacher (Rahul)', P, 'Today', 'Log in as Rahul.', 'rahul@demo.test', 'Std 5 A card shows attendance already done (green). "My subjects" lists Maths and English for Std 5 A, 6 A, 7 A.', 'Medium'),
    # ---- Office clerk
    ('Office clerk', 'Office clerk', C, 'Menu', 'Log in as the clerk.', 'clerk@demo.test', 'Opens the Dashboard. Menu has Dashboard, Attendance, Students, Syllabus, Me. No Marks and no school setup pages.', 'High'),
    ('Office clerk', 'Office clerk', C, 'Attendance → Std 1 A', 'Open Attendance, then a class.', '', 'The class list is shown with a note "You can see this attendance, but only teachers and the principal can change it." No Save button.', 'High'),
    ('Office clerk', 'Office clerk', C, 'Students → Add student', 'Add a child with all details: open "More details (parents, address, documents)" and fill them. Leave "Register number (GR)" empty.', 'QA Child Two, Std 2 A, Boy, DOB 10/05/2019, parent Sunita Rane 9876501234, address Kalwa', 'Saved with the next register number (SS-…). All details show on the profile.', 'High'),
    ('Office clerk', 'Office clerk', C, 'Students → Add student', 'Tick "Date of birth not known" and fill approximate age.', 'QA Child Three, Std 3 A, age 8', 'Saved. The profile shows the age as approximate.', 'Medium'),
    ('Office clerk', 'Office clerk', C, 'Students → Add student', 'Type a wrong mobile number.', '12345', 'Message "Enter a valid mobile number". Not saved until fixed.', 'Medium'),
    ('Office clerk – Import', 'Office clerk', C, 'Students → Import from Excel', 'Tap "Download template" and open the file.', '', 'An Excel file downloads with columns name*, class*, gender, date_of_birth … and a Help sheet listing the classes.', 'Medium'),
    ('Office clerk – Import', 'Office clerk', C, 'Students → Import from Excel', 'Choose the good file. Check the rows. Tap "Import 5 students".', 'test-files/students-import-5-good.xlsx', '"5 of 5 rows are ready". After import: "5 students imported". Asha Lokhande and Rohan Dhotre are now in Std 1 A.', 'High'),
    ('Office clerk – Import', 'Office clerk', C, 'Students → Import from Excel', 'Choose the file with mistakes.', 'test-files/students-import-with-errors.xlsx', '"1 of 5 rows are ready". Problems in red: Ajay Borse "Class: Class not found", row 4 "Full name: Required", Tara Jogdand "Date of birth: Wrong date (use DD/MM/YYYY)", Vikas Ingle "Parent mobile number: Please check this". "Import 1 student" adds only Meghna Salve.', 'High'),
    ('Office clerk – Import', 'Office clerk', P, 'Students → Import from Excel', 'On a phone, import the CSV file.', 'test-files/students-import-3-rows.csv', '"3 of 3 rows are ready", then "3 students imported".', 'Medium'),
    ('Office clerk – Import', 'Office clerk', C, 'Students → Import from Excel', 'Choose the PDF file.', 'test-files/not-a-photo.pdf', '"This file type is not allowed". Nothing imported.', 'Medium'),
    ('Office clerk – Leaving', 'Office clerk', B, 'Nikhil Sonawane → Mark as left school', 'Open Nikhil Sonawane. Tap "Mark as left school". Choose a reason, add a note, confirm.', 'Reason: Joined another school; note: Moved to TMC School No. 12', 'Status "Left school" with the date. Nikhil is no longer in the Std 3 A attendance list.', 'High'),
    ('Office clerk – Leaving', 'Office clerk', B, 'Nikhil Sonawane → Leaving certificate', 'Tap "Leaving certificate", then "Print / Save PDF".', '', 'Certificate with school name, GR number, name, parents, date of birth, last class, attendance, date of leaving. Print / save as PDF opens; nothing is cut off.', 'High'),
    ('Office clerk – Leaving', 'Office clerk', C, 'Riya Kale → (address)/certificate', 'Open Riya Kale, add /certificate at the end of the address.', '', 'Message "This child is still studying here. Mark the child as left before printing the certificate."', 'Low'),
    ('Office clerk – Search', 'Office clerk', B, 'Students → search', 'Search "Lata". Then tick "Search all years".', 'Lata', 'Without the tick: "No students found". With the tick: Lata Bhalerao, "Left school".', 'High'),
    ('Office clerk – Search', 'Office clerk', B, 'Lata Bhalerao → Re-admit', 'Open Lata. Tap "Re-admit", choose Std 3 A, confirm.', 'Std 3 A', '"Saved". Lata is "Studying" in Std 3 A and appears in Std 3 A attendance. Her old year stays under "Past years".', 'High'),
    ('Office clerk – Health', 'Office clerk', B, 'Swati Gholap → Health', 'Open Swati Gholap, tab Health.', '', 'Message "The last check-up says this child needs to see a doctor again." with the note about her eyes.', 'Medium'),
    ('Office clerk – Students', 'Office clerk', C, 'Students → filters', 'Use the Class filter (Std 4 A), then Status "Left school", then "All".', '', 'The list and the count ("… children") change with each filter.', 'Medium'),
    ('Office clerk – Students', 'Office clerk', C, 'Students → Download Excel', 'Tap "Download Excel".', '', 'An Excel file with the children of the list downloads and opens.', 'Low'),
    ('Office clerk – Register', 'Office clerk', C, 'Attendance → Monthly register', 'Choose Std 1 A and this month. Tap "Download Excel", then "Print / Save PDF".', '', 'Grid of children × days with P / A / L / LT, grey holidays, "Present each day" row. Excel downloads; print shows the whole register.', 'Medium'),
    # ---- Principal
    ('Principal – Dashboard', 'Principal', B, 'Dashboard', 'Log in as the principal.', 'principal@demo.test', 'Dashboard with boxes: Students, Teachers, Classes, Attendance done today (x/8), Absent today, Meals to serve today. Numbers are not empty.', 'High'),
    ('Principal – Dashboard', 'Principal', B, 'Dashboard', 'Look at "Attendance not taken".', '', 'Lists classes still waiting today (at the start: Std 1 A and Std 3 A). After a teacher saves, that class goes away.', 'High'),
    ('Principal – Dashboard', 'Principal', P, 'Dashboard → Absent several days in a row', 'Find Kiran Thakur. Tap her WhatsApp icon.', 'Kiran Thakur', 'Kiran Thakur, Std 4 A, "5 days". WhatsApp opens with a ready message to the family.', 'High'),
    ('Principal – Dashboard', 'Principal', B, 'Dashboard → Children who often miss school', 'Look at the list.', '', 'Kiran Thakur is listed with her % (below 75%) and Std 4 A.', 'Medium'),
    ('Principal – Dashboard', 'Principal', B, 'Dashboard (bottom)', 'Scroll down.', '', '"Birthdays today: Ishaan Jagtap (Std 2 A)". Also "Attendance in the last 2 weeks" bars and "Syllabus progress" per class.', 'Medium'),
    ('Principal – Dashboard', 'Principal', B, 'Dashboard', 'Look for the health box.', '', '"… children need to see a doctor again (from the last health camp)".', 'Low'),
    ('Principal – Attendance', 'Principal', B, 'Attendance', 'Open Attendance. Tap a class that is done. Change one mark and save.', '', 'Classes show "x present · y absent" or "Not done". The principal can change and save.', 'High'),
    ('Principal – Attendance', 'Principal', B, 'Attendance → class → Date', 'Pick a date 10 days ago and change one mark. Save.', '10 days ago', 'Saved (the principal may correct old days; a teacher may not).', 'Medium'),
    ('Principal – Staff', 'Principal', C, 'More → Teachers & staff → Add teacher or staff', 'Add a new teacher.', 'Name: QA Teacher Two; Mobile 9820000099; Role Teacher', 'Box "Login details" with a temporary password. "Copy message for WhatsApp" → "Copied…". The teacher is in the Active list.', 'High'),
    ('Principal – Staff', 'Principal', C, 'Teachers & staff → Add', 'Try to add someone with a mobile number already used.', 'Mobile 9820000004', 'Not added: message "This already exists" (or "Already used" under the box).', 'Medium'),
    ('Principal – Staff', 'Principal', C, 'Teachers & staff → Neha Joshi → Classes', 'Open Neha Joshi, tab "Classes". Add: Std 2 A, Subject teacher, subject English. Then log in as Neha.', 'newteacher@demo.test / NehaQa@2026', 'The class shows in her list. Neha\'s Today page now shows "English · Std 2 A" under My subjects.', 'High'),
    ('Principal – Staff', 'Principal', C, 'Neha Joshi → Reset password', 'Tap "Reset password" and confirm. Log in as Neha with the new temporary password.', '', 'A new temporary password is shown once. Neha\'s old password stops working; with the new one she must choose her own password again.', 'High'),
    ('Principal – Staff', 'Principal', C, 'QA Teacher Two → Make inactive', 'Tap "Make inactive" and confirm. Try to log in as that teacher. Then "Make active again".', '', 'Teacher moves to the "Inactive" tab and cannot log in ("This account is inactive…"). After "Make active again" login works.', 'Medium'),
    ('Principal – Classes', 'Principal', C, 'More → Classes & subjects', 'In Std 4 A type "Computer" in "New subject" and tap Add. Then remove it (x) and confirm.', 'Computer', 'The subject appears as a chip and then is removed.', 'Medium'),
    ('Principal – Classes', 'Principal', C, 'Classes & subjects → Class levels', 'Open "Class levels".', '', 'Levels Balwadi to Std 7 with their "Next class level"; Std 7 is "Last class of the school".', 'Low'),
    ('Principal – Classes', 'Principal', C, 'Classes & subjects → Activities', 'Add a new activity. Open its members, tick 2 children and save.', 'Chess club; Kiran Thakur, Riya Kale', 'The activity shows "2 members".', 'Low'),
    ('Principal – Syllabus', 'Principal', C, 'Syllabus → a subject → Edit chapters', 'Add a chapter, paste 3 topics (one per line), tap "Add these topics". Remove one new topic. Try to remove a topic that is already taught.', 'QA Chapter; Topic A / Topic B / Topic C', 'Chapter with 3 topics appears; removing a new topic works; a taught topic cannot be removed ("Already taught, so it cannot be removed").', 'Medium'),
    ('Principal – Holidays', 'Principal', C, 'More → Holidays', 'Add a holiday for a date next week. Then delete it.', 'QA Holiday, next week, type Holiday', 'Holiday shows in the list, then is removed. (Do NOT add a holiday for today.)', 'Medium'),
    ('Principal – Years', 'Principal', C, 'More → Academic years', 'Look at the list.', '', '2026-27 "Current", older years "Closed", each with dates, classes and number of students (not 0).', 'Medium'),
    ('Principal – Years', 'Principal', C, 'Academic years → Start new academic year', 'Go through all 5 steps for 2027-28. On the last step keep "Make it current now" OFF. Tap "Create the new year".', '2027-28', '"2027-28 is ready" with counts. The list shows 2027-28 as "Planned". IMPORTANT: do NOT make it current.', 'High'),
    ('Principal – Years', 'Principal', C, 'Academic years → (closed year) → Unlock for corrections', 'Unlock the last closed year for 10 minutes with a reason.', '10 minutes; reason: QA test', '"Unlocked until (time)". The activity log shows "Unlocked a year".', 'Low'),
    ('Principal – School', 'Principal', C, 'More → School settings → This school', 'Change the phone number and save. Tap "Upload school logo" and choose the logo.', 'Phone 022 2500 0099; test-files/school-logo.png', '"Saved". The new logo appears at the top left of the app.', 'Medium'),
    ('Principal – School', 'Principal', C, 'School settings → This school', 'Look at "Register number prefix" and "Days teachers can change old attendance".', '', 'Help text shows the next register number (e.g. "Next register number will be SS-…"). Do not change weekly holidays.', 'Low'),
    ('Principal – Activity log', 'Principal', C, 'More → Activity log', 'Filter Person = Sunita Patil, then Kind of change = Attendance. Use "Older" / "Newer".', '', 'Shows who changed what and when, e.g. "Sunita Patil · Saved attendance · Std 1 A". Filters and pages work.', 'Medium'),
    ('Principal – Phone', 'Principal', P, 'All pages on a phone', 'On a phone, open every page from the bottom bar and "More".', '', 'Every page fits the screen (no sideways scroll), all buttons can be tapped, nothing is hidden behind the bottom bar.', 'High'),
    # ---- Owner
    ('Trust owner', 'Trust owner', B, 'Me → School (or the school name)', 'Switch to "Signal Shala, Nerul". Then switch back to Thane.', 'owner@demo.test', 'Nerul data is shown and the app colours change to the Marigold look; back in Thane the red Signal Shala look returns.', 'High'),
    ('Trust owner', 'Trust owner', C, 'School settings → All schools', 'Look at the table. Add a school.', 'QA School, Kalyan', 'Both schools with students, teachers and attendance. The new school is added to the table.', 'Low'),
    ('Trust owner – second school', 'Teacher (2nd school)', P, 'Today', 'Log in as Imran.', 'imran@demo.test', 'Only Signal Shala, Nerul is shown, in the Marigold look. Thane children are not visible.', 'Medium'),
    # ---- Looks (themes)
    ('Looks', 'Principal', B, 'More → School settings → Look', 'Open the Look tab.', '', '6 looks with a small preview each: Signal Shala (In use), Blackboard, Notebook, Marigold, Calm teal, Easy to read.', 'High'),
    ('Looks', 'Principal', B, 'School settings → Look', 'Tap "Blackboard". Look around. Tap "Cancel".', '', '"You are trying “Blackboard”. Only you can see it." Colours, letters and icons change at once. Cancel returns to Signal Shala.', 'High'),
    ('Looks', 'Principal + Teacher', B, 'School settings → Look', 'Tap "Notebook", then "Use for the whole school". On another phone, log in (or reload) as Sunita.', 'sunita@demo.test', '"Everyone in the school now sees “Notebook”". Sunita\'s app is in the Notebook look too.', 'High'),
    ('Looks', 'Principal', B, 'Look: Signal Shala', 'Use this look for the school. Open Today, Take attendance, Students, a student, Dashboard (as principal) on a phone.', '', 'Red/white look with the traffic-light stripe. All text readable, buttons clear, icons match, nothing overlaps.', 'High'),
    ('Looks', 'Principal', B, 'Look: Blackboard', 'Same pages as above.', '', 'Green-board look with chalk-yellow and handwritten headings. Everything readable, no overlap.', 'Medium'),
    ('Looks', 'Principal', B, 'Look: Notebook', 'Same pages as above.', '', 'Blue ink, ruled-page cards, sharp corners. Everything readable, no overlap.', 'Medium'),
    ('Looks', 'Principal', B, 'Look: Marigold', 'Same pages as above.', '', 'Saffron and maroon, soft round shapes. Everything readable, no overlap.', 'Medium'),
    ('Looks', 'Principal', B, 'Look: Calm teal', 'Same pages as above.', '', 'Teal and white, simple. Everything readable, no overlap.', 'Medium'),
    ('Looks', 'Principal', B, 'Look: Easy to read', 'Same pages as above, also outside in sunlight if possible.', '', 'Black on white, thick borders, bigger letters. Very easy to read.', 'Medium'),
    ('Looks', 'Teacher', P, 'More → Me', 'As Sunita, look for a Look/theme setting.', '', 'Teachers cannot change the school look (no Look option).', 'Medium'),
    ('Looks', 'Any', P, 'Phone in dark mode', 'Turn on Dark mode in the phone settings. Open the app.', '', 'The app stays light (no dark mode in any look).', 'Low'),
    ('Looks', 'Principal', B, 'School settings → Look', 'At the end of testing: choose "Signal Shala" → "Use for the whole school".', '', 'The school is back on the Signal Shala look.', 'High'),
    # ---- General
    ('General', 'Any', P, 'Every page', 'Use a small phone (or the smallest screen you have). Open each page.', '', 'No page scrolls sideways; text is not cut; buttons are big enough to tap with a thumb.', 'High'),
    ('General', 'Any', P, 'Any page', 'Turn the phone sideways (landscape) on Today and Take attendance.', '', 'Pages still work and look fine.', 'Low'),
    ('General', 'Any', P, 'Chrome menu → Install app / Add to Home screen', 'Install the app and open it from the home screen.', '', 'Opens like an app (no browser address bar), with the Signal School icon.', 'Medium'),
    ('General', 'Any', P, '? buttons', 'Tap the "?" on Today, Students, Syllabus, Academic years.', '', 'A short help text in simple words appears.', 'Low'),
    ('General', 'Any', P, 'Language check (मराठी / ગુજરાતી)', 'Switch to मराठी and then ગુજરાતી. Open Today, Attendance, Students, a student, Me.', '', 'All words in that language and easy to read. Write in Comments any English words left or wrong spellings.', 'Medium'),
    ('General', 'Any', B, 'Phone Back button / browser Back', 'Move through several pages and press Back.', '', 'Back goes to the previous page every time (does not exit the app unexpectedly).', 'Medium'),
    ('General', 'Any', C, 'Computer, big window', 'Open on a computer. Make the window narrow, then wide.', '', 'Wide: side menu with groups "Every day", "School setup", "Other". Narrow: bottom bar like the phone.', 'Medium'),
    ('General', 'Any', B, 'Slow internet', 'In Chrome, use a slow connection (or weak mobile signal). Open Today and Students.', '', 'Pages show a loading state and then load; nothing breaks.', 'Low'),
]


def build():
    wb = Workbook()
    thin = Side(style='thin', color='D9D9D9')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill('solid', fgColor=RED)
    head_font = Font(bold=True, color='FFFFFF')
    wrap = Alignment(wrap_text=True, vertical='top')

    def header(ws, cols, widths):
        ws.append(cols)
        for i, w in enumerate(widths, 1):
            c = ws.cell(row=1, column=i)
            c.fill, c.font, c.border = head_fill, head_font, border
            c.alignment = Alignment(wrap_text=True, vertical='center')
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'A2'
        ws.row_dimensions[1].height = 30

    def rows(ws, data):
        for r in data:
            ws.append(list(r))
            for c in ws[ws.max_row]:
                c.alignment, c.border = wrap, border

    # Start here
    ws = wb.active
    ws.title = 'Start here'
    ws.column_dimensions['A'].width = 110
    lines = [
        ('Signal School – manual testing', Font(bold=True, size=16, color=RED)),
        ('Read the PDF guide first (Signal-School-QA-Guide.pdf). It explains the school, the app and how to test.', None),
        ('', None),
        ('How to use this file', Font(bold=True, size=12)),
        ('1. Go to the sheet "Test cases". Do the tests from top to bottom (some tests use what an earlier test did).', None),
        ('2. For each row: log in as the person in "Who logs in" (see sheet "Logins"), follow "Steps", use "Test data".', None),
        ('3. Compare what you see with "Expected result". Choose a Status: Pass, Fail, Blocked or Skipped.', None),
        ('4. Write anything odd in "Comments" – even small things (a word cut off, a slow page, a confusing button).', None),
        ('5. For every Fail, add a row in the sheet "Bugs" with a screenshot name. Put the screenshots in one folder.', None),
        ('6. The sheet "Summary" counts Pass / Fail automatically.', None),
        ('', None),
        ('Status meaning', Font(bold=True, size=12)),
        ('Pass – works as written.   Fail – something is different or broken.   Blocked – could not test (e.g. earlier step failed).', None),
        ('Skipped – not tested on purpose (write why).   Not run – not done yet.', None),
        ('', None),
        ('Good to know', Font(bold=True, size=12)),
        ('• "Phone" tests: use a real Android phone with Chrome. "Computer" tests: Chrome on a laptop/desktop. "Phone + Computer": do both.', None),
        ('• The files to upload (photos, Excel sheets) are in the folder "test-files" next to this file. Copy them to the phone first.', None),
        ('• The children and people are made-up test data. Change and add things freely.', None),
        ('• If the data gets mixed up and tests cannot continue, ask for a "test data reset" (everything goes back to the start).', None),
        ('• Never use real children\'s names, photos or phone numbers while testing.', None),
    ]
    for text, font in lines:
        ws.append([text])
        cell = ws.cell(row=ws.max_row, column=1)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        if font:
            cell.font = font

    # Test cases
    ws = wb.create_sheet('Test cases')
    cols = ['ID', 'Area', 'Who logs in', 'Device', 'Where (screen / path)', 'Steps', 'Test data', 'Expected result', 'Priority', 'Status', 'Comments', 'Tested on (date)', 'Tester']
    header(ws, cols, [8, 18, 16, 12, 26, 48, 26, 52, 9, 11, 34, 13, 12])
    prefix = {}
    data = []
    for area, who, device, where, steps, tdata, expected, prio in CASES:
        key = area.split(' – ')[0]
        code = {'Login': 'LG', 'Teacher': 'TC', 'Office clerk': 'CL', 'Principal': 'PR', 'Trust owner': 'OW', 'Looks': 'LK', 'General': 'GN'}[key]
        prefix[code] = prefix.get(code, 0) + 1
        data.append((f'{code}-{prefix[code]:02d}', area, who, device, where, steps, tdata, expected, prio, 'Not run', '', '', ''))
    rows(ws, data)
    n = len(data) + 1
    dv = DataValidation(type='list', formula1='"' + ','.join(STATUSES) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f'J2:J{n}')
    for status, color in [('Pass', 'C6EFCE'), ('Fail', 'FFC7CE'), ('Blocked', 'FFEB9C'), ('Skipped', 'E7E6E6')]:
        ws.conditional_formatting.add(f'J2:J{n}', FormulaRule(formula=[f'$J2="{status}"'], fill=PatternFill('solid', fgColor=color)))
    ws.conditional_formatting.add(f'I2:I{n}', FormulaRule(formula=['$I2="High"'], font=Font(bold=True, color=RED)))
    ws.auto_filter.ref = f'A1:M{n}'

    # Logins
    ws = wb.create_sheet('Logins')
    header(ws, ['Role', 'Name', 'Email', 'Mobile', 'Password', 'What to use it for'], [18, 18, 24, 14, 24, 70])
    rows(ws, LOGINS)

    # Test data
    ws = wb.create_sheet('Test data')
    header(ws, ['Child', 'Class', 'About', 'Used for'], [18, 20, 60, 50])
    rows(ws, CHILDREN)
    ws.append([])
    ws.append(['File (folder test-files)', 'What it is', 'Used for'])
    for c in ws[ws.max_row]:
        c.fill, c.font = head_fill, head_font
    rows(ws, FILES)
    ws.append([])
    ws.append(['Schools: Signal Shala, Thane (Balwadi to Std 7, look "Signal Shala") and Signal Shala, Nerul (Balwadi to Std 2, look "Marigold").'])
    ws.append(['Every day counts as a school day in the test data, so attendance can be tested any day of the week.'])

    # Bugs
    ws = wb.create_sheet('Bugs')
    header(ws, ['Bug no.', 'Test case ID', 'Screen', 'What happened', 'What should happen', 'Steps to see it again', 'Phone / Computer (model, browser)', 'Screenshot file', 'How bad (Blocker / Major / Minor / Cosmetic)', 'Date', 'Fixed? (for developer)'],
           [8, 12, 20, 40, 34, 40, 22, 18, 16, 12, 14])
    for i in range(1, 41):
        ws.append([i] + [''] * 10)
        for c in ws[ws.max_row]:
            c.alignment, c.border = wrap, border
    sev = DataValidation(type='list', formula1='"Blocker,Major,Minor,Cosmetic"', allow_blank=True)
    ws.add_data_validation(sev)
    sev.add('I2:I41')

    # Summary
    ws = wb.create_sheet('Summary')
    ws.column_dimensions['A'].width = 26
    for col in 'BCDEFG':
        ws.column_dimensions[col].width = 11
    ws.append(['Area', 'Total', 'Pass', 'Fail', 'Blocked', 'Skipped', 'Not run'])
    for c in ws[1]:
        c.fill, c.font, c.border = head_fill, head_font, border
    groups = []
    for area in [c[0] for c in CASES]:
        g = area.split(' – ')[0]
        if g not in groups:
            groups.append(g)
    rng = f"'Test cases'!$B$2:$B${n}"
    srng = f"'Test cases'!$J$2:$J${n}"
    for g in groups:
        r = ws.max_row + 1
        ws.append([g, f'=COUNTIF({rng},"{g}*")'] + [f'=COUNTIFS({rng},"{g}*",{srng},"{s}")' for s in ['Pass', 'Fail', 'Blocked', 'Skipped', 'Not run']])
    ws.append(['All tests', f'=COUNTA({srng})'] + [f'=COUNTIF({srng},"{s}")' for s in ['Pass', 'Fail', 'Blocked', 'Skipped', 'Not run']])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.border = border

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f'{OUT}: {len(CASES)} test cases')


if __name__ == '__main__':
    build()
